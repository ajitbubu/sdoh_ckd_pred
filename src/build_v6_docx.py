"""
src/build_v6_docx.py
--------------------
Builds the JMIR PHS v6 resubmission from docs/CKD_Paper_JMIR_PHS_v6_revised.md:

  docs/CKD_Paper_JMIR_PHS_v6_revised.docx              clean version
  docs/CKD_Paper_JMIR_PHS_v6_revised_highlighted.docx  new/changed text vs v5
                                                       highlighted in yellow
  docs/CKD_Paper_JMIR_PHS_v6_Multimedia_Appendix_1.docx  supplementary tables,
                                                       split out of the manuscript
  docs/CKD_Paper_JMIR_PHS_v6_Multimedia_Appendix_3_TRIPOD-AI.docx  TRIPOD+AI checklist
  docs/Multimedia_Appendix_2_pipeline_metrics.zip      metric files from
                                                       ckd_pipeline/outputs/

Citations are written as {{key}} in the Markdown and numbered in order of
first appearance (JMIR reference style). Styles come from
docs/reference_styles.docx (the v5 styles, no content) via pandoc --reference-doc.
The highlighted version needs the previous (v5) docx, which is not in the
repository, and is skipped without it. Requires pandoc and python-docx.
"""

import difflib
import os
import re
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(BASE_DIR, "docs")
MD = os.path.join(DOCS, "CKD_Paper_JMIR_PHS_v6_revised.md")
V5 = os.path.join(DOCS, "CKD_Paper_JMIR_PHS_v5_revised.docx")  # previous version; local only
# Styles only (from the v5 docx, no content), so the build runs from a fresh clone
REFERENCE = os.path.join(DOCS, "reference_styles.docx")
CLEAN = os.path.join(DOCS, "CKD_Paper_JMIR_PHS_v6_revised.docx")
HIGHLIGHTED = os.path.join(DOCS, "CKD_Paper_JMIR_PHS_v6_revised_highlighted.docx")
APPENDIX_1 = os.path.join(DOCS, "CKD_Paper_JMIR_PHS_v6_Multimedia_Appendix_1.docx")
APPENDIX_3_MD = os.path.join(DOCS, "CKD_Paper_JMIR_PHS_v6_Multimedia_Appendix_3_TRIPOD-AI.md")
APPENDIX_3 = APPENDIX_3_MD.replace(".md", ".docx")
APPENDIX_2 = os.path.join(DOCS, "Multimedia_Appendix_2_pipeline_metrics.zip")
OUTPUTS = os.path.join(BASE_DIR, "ckd_pipeline", "outputs")

REFS = {
    "cdc2023": "Centers for Disease Control and Prevention. Chronic Kidney Disease in the United States, 2023. Atlanta, GA: US Department of Health and Human Services; 2023. URL: https://archive.cdc.gov/www_cdc_gov/kidney-disease/php/data-research/index.html [accessed 2026-10-24]",
    "usrds2023": "United States Renal Data System. 2023 USRDS Annual Data Report: Epidemiology of Kidney Disease in the United States. Bethesda, MD: National Institutes of Health, National Institute of Diabetes and Digestive and Kidney Diseases; 2023. URL: https://usrds-adr.niddk.nih.gov/2023",
    "vart2015": "Vart P, Gansevoort RT, Joosten MM, Bültmann U, Reijneveld SA. Socioeconomic disparities in chronic kidney disease: a systematic review and meta-analysis. Am J Prev Med. 2015;48(5):580-592. PMID: 25891058. doi:10.1016/j.amepre.2014.11.004",
    "crews2010": "Crews DC, Charles RF, Evans MK, Zonderman AB, Powe NR. Poverty, race, and CKD in a racially and socioeconomically diverse urban population. Am J Kidney Dis. 2010;55(6):992-1000. PMID: 20207457. doi:10.1053/j.ajkd.2009.12.032",
    "nicholas2015": "Nicholas SB, Kalantar-Zadeh K, Norris KC. Socioeconomic disparities in chronic kidney disease. Adv Chronic Kidney Dis. 2015;22(1):6-15. PMID: 25573507. doi:10.1053/j.ackd.2014.07.002",
    "alderwick2019": "Alderwick H, Gottlieb LM. Meanings and misunderstandings: a social determinants of health lexicon for health care systems. Milbank Q. 2019;97(2):407-419. PMID: 31069864. doi:10.1111/1468-0009.12390",
    "places_notes": "Centers for Disease Control and Prevention. PLACES: Local Data for Better Health — current release notes. URL: https://www.cdc.gov/places/current-release-notes/ [accessed 2026-10-17]",
    "bang2007": "Bang H, Vupputuri S, Shoham DA, et al. SCreening for Occult REnal Disease (SCORED): a simple prediction model for chronic kidney disease. Arch Intern Med. 2007;167(4):374-381. PMID: 17325299. doi:10.1001/archinte.167.4.374",
    "kshirsagar2008": "Kshirsagar AV, Bang H, Bomback AS, et al. A simple algorithm to predict incident kidney disease. Arch Intern Med. 2008;168(22):2466-2473. PMID: 19064831. doi:10.1001/archinte.168.22.2466",
    "sanmarchi2023": "Sanmarchi F, Fanconi C, Golinelli D, Gori D, Hernandez-Boussard T, Capodici A. Predict, diagnose, and treat chronic kidney disease with machine learning: a systematic literature review. J Nephrol. 2023;36(4):1101-1117. PMID: 36786976. doi:10.1007/s40620-023-01573-4",
    "tangri2011": "Tangri N, Stevens LA, Griffith J, et al. A predictive model for progression of chronic kidney disease to kidney failure. JAMA. 2011;305(15):1553-1559. PMID: 21482743. doi:10.1001/jama.2011.451",
    "adi2020": "University of Wisconsin School of Medicine and Public Health, Center for Health Disparities Research. Area Deprivation Index, 2020, version 4.0.1, national census block group file. Neighborhood Atlas. URL: https://www.neighborhoodatlas.medicine.wisc.edu/ (free registration and data use agreement required) [accessed 2026-10-17]",
    "kind2018": "Kind AJH, Buckingham WR. Making neighborhood-disadvantage metrics accessible — the Neighborhood Atlas. N Engl J Med. 2018;378(26):2456-2458. PMID: 29949490. doi:10.1056/NEJMp1802313",
    "places2022": "Centers for Disease Control and Prevention, Division of Population Health. PLACES: Census Tract Data (GIS Friendly Format), 2022 release. Socrata dataset shc3-fzig. URL: https://data.cdc.gov/500-Cities-Places/PLACES-Census-Tract-Data-GIS-Friendly-Format-2022-/shc3-fzig. Data file: https://data.cdc.gov/api/views/shc3-fzig/rows.csv?accessType=DOWNLOAD [accessed 2026-10-25]",
    "zhang2014": "Zhang X, Holt JB, Lu H, et al. Multilevel regression and poststratification for small-area estimation of population health outcomes: a case study of chronic obstructive pulmonary disease prevalence using the Behavioral Risk Factor Surveillance System. Am J Epidemiol. 2014;179(8):1025-1033. PMID: 24598867. doi:10.1093/aje/kwu018",
    "nhanes_2123_guidance": "National Center for Health Statistics. NHANES August 2021–August 2023: brief overview and analytic guidance. URL: https://wwwn.cdc.gov/nchs/nhanes/ContinuousNhanes/OverviewBrief.aspx?Cycle=2021-2023 [accessed 2026-10-24]",
    "nhanes": "National Center for Health Statistics. National Health and Nutrition Examination Survey (NHANES), 2017–March 2020 Pre-Pandemic Combined Cycle and August 2021–August 2023 Cycle. Public-use data files. URL: https://wwwn.cdc.gov/nchs/nhanes/. Data files: https://wwwn.cdc.gov/Nchs/Data/Nhanes/Public/2017/DataFiles/ and https://wwwn.cdc.gov/Nchs/Data/Nhanes/Public/2021/DataFiles/ [accessed 2026-10-25]",
    "xgboost": "Chen T, Guestrin C. XGBoost: a scalable tree boosting system. In: Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining. New York, NY: ACM; 2016:785-794. doi:10.1145/2939672.2939785",
    "kdigo2024": "Kidney Disease: Improving Global Outcomes (KDIGO) CKD Work Group. KDIGO 2024 clinical practice guideline for the evaluation and management of chronic kidney disease. Kidney Int. 2024;105(4S):S117-S314. PMID: 38490803. doi:10.1016/j.kint.2023.10.018",
    "inker2021": "Inker LA, Eneanya ND, Coresh J, et al. New creatinine- and cystatin C–based equations to estimate GFR without race. N Engl J Med. 2021;385(19):1737-1749. PMID: 34554658. doi:10.1056/NEJMoa2102953",
    "murphy2016": "Murphy D, McCulloch CE, Lin F, et al; Centers for Disease Control and Prevention Chronic Kidney Disease Surveillance Team. Trends in prevalence of chronic kidney disease in the United States. Ann Intern Med. 2016;165(7):473-481. PMID: 27479614. doi:10.7326/M16-0273",
    "coresh2007": "Coresh J, Selvin E, Stevens LA, et al. Prevalence of chronic kidney disease in the United States. JAMA. 2007;298(17):2038-2047. PMID: 17986697. doi:10.1001/jama.298.17.2038",
    "vancalster2019": "Van Calster B, McLernon DJ, van Smeden M, et al; Topic Group 'Evaluating diagnostic tests and prediction models' of the STRATOS initiative. Calibration: the Achilles heel of predictive analytics. BMC Med. 2019;17(1):230. PMID: 31842878. doi:10.1186/s12916-019-1466-7",
    "vickers2006": "Vickers AJ, Elkin EB. Decision curve analysis: a novel method for evaluating prediction models. Med Decis Making. 2006;26(6):565-574. PMID: 17099194. doi:10.1177/0272989X06295361",
    "raowu1988": "Rao JNK, Wu CFJ. Resampling inference with complex survey data. J Am Stat Assoc. 1988;83(401):231-241. doi:10.1080/01621459.1988.10478591",
    "github": "Sahu A, Mukherjee S, Poreddy KKR. SDOH-CKDPred: open CKD small-area surveillance and screening-triage pipeline, version 6.0.1. Zenodo [AUTHORS: insert DOI]. Source repository URL: https://github.com/ajitbubu/sdoh_ckd_pred",
    "adi_changelog": "University of Wisconsin School of Medicine and Public Health, Center for Health Disparities Research. Neighborhood Atlas: changes between versions of the ADI. URL: https://www.neighborhoodatlas.medicine.wisc.edu/changelog [accessed 2026-10-17]",
    "austin2019": "Austin PC, Steyerberg EW. The Integrated Calibration Index (ICI) and related metrics for quantifying the calibration of logistic regression models. Stat Med. 2019;38(21):4051-4065. PMID: 31270850. doi:10.1002/sim.8281",
    "nchs_erb": "National Center for Health Statistics. NHANES — NCHS Ethics Review Board approval. URL: https://www.cdc.gov/nchs/nhanes/about/erb.html",
    "cfr46": "Protection of Human Subjects. 45 CFR §46.102 (2018). URL: https://www.ecfr.gov/current/title-45/subtitle-A/subchapter-A/part-46",
    "wakefield2020": "Wakefield J, Okonek T, Pedersen J. Small area estimation for disease prevalence mapping. Int Stat Rev. 2020;88(2):398-418. PMID: 36081593. doi:10.1111/insr.12400",
    "collins2024": "Collins GS, Moons KGM, Dhiman P, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378. PMID: 38626948. doi:10.1136/bmj-2023-078378",
    "christodoulou2019": "Christodoulou E, Ma J, Collins GS, Steyerberg EW, Verbakel JY, Van Calster B. A systematic review shows no performance benefit of machine learning over logistic regression for clinical prediction models. J Clin Epidemiol. 2019;110:12-22. PMID: 30763612. doi:10.1016/j.jclinepi.2019.02.004",
    "delgado2022": "Delgado C, Baweja M, Crews DC, et al. A unifying approach for GFR estimation: recommendations of the NKF-ASN Task Force on Reassessing the Inclusion of Race in Diagnosing Kidney Disease. Am J Kidney Dis. 2022;79(2):268-288.e1. PMID: 34563581. doi:10.1053/j.ajkd.2021.08.003",
    "census_crosswalk": "US Census Bureau. 2020 Census Tract to 2010 Census Tract Relationship File. Geography Division; 2021. URL: https://www.census.gov/geographies/reference-files/time-series/geo/relationship-files.2020.html",
}


def jmir_style(ref):
    """'... PMID: n. doi:x' -> '... [doi: x] [Medline: n]' (JMIR reference style)."""
    pmid = re.search(r"\s*PMID: (\d+)\.", ref)
    doi = re.search(r"\s*(?:doi:|DOI:) ?(\S+?)\.?(?=\.? Source|$)", ref)
    ids = []
    if doi:
        ids.append(f"[doi: {doi.group(1)}]")
    if pmid:
        ids.append(f"[Medline: {pmid.group(1)}]")
    for m in (pmid, doi):
        if m:
            ref = ref.replace(m.group(0), "", 1)
    ref = ref.rstrip(" .") + "."
    return f"{ref} {' '.join(ids)}" if ids else ref


def number_citations(text):
    order = []
    for key in re.findall(r"\{\{(\w+)\}\}", text):
        if key not in ("REFERENCES",) and key not in order:
            order.append(key)
    missing = [k for k in order if k not in REFS]
    assert not missing, f"unknown citation keys: {missing}"
    num = {k: i + 1 for i, k in enumerate(order)}

    def group(m):
        ns = sorted(num[k] for k in re.findall(r"\{\{(\w+)\}\}", m.group(0)))
        runs = [[ns[0]]]
        for n in ns[1:]:
            if n == runs[-1][-1] + 1:
                runs[-1].append(n)
            else:
                runs.append([n])
        return "[" + ",".join(f"{r[0]}-{r[-1]}" if len(r) > 2 else
                              ",".join(map(str, r)) for r in runs) + "]"

    text = re.sub(r"(?:\{\{(?!REFERENCES)\w+\}\})+", group, text)
    refs = "\n\n".join(f"{num[k]}. {jmir_style(REFS[k])}" for k in order)
    unused = set(REFS) - set(order)
    if unused:
        print(f"  note: unused references {sorted(unused)}")
    return text.replace("{{REFERENCES}}", refs)


def norm(s):
    return re.sub(r"\s+", " ", s).strip().lower()


def highlight_changes(clean_path, v5_path, out_path):
    old = [norm(p.text) for p in Document(v5_path).paragraphs if p.text.strip()]
    old += [norm(c.text) for t in Document(v5_path).tables
            for row in t.rows for c in row.cells if c.text.strip()]

    def unchanged(text):
        t = norm(text)
        if t in old:
            return True
        best = difflib.get_close_matches(t, old, n=1, cutoff=0.9)
        return bool(best)

    doc = Document(clean_path)
    n_changed = 0
    paras = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows
                                    for c in row.cells for p in c.paragraphs]
    for p in paras:
        if p.text.strip() and not unchanged(p.text):
            n_changed += 1
            for r in p.runs:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.save(out_path)
    return n_changed, len([p for p in paras if p.text.strip()])


def split_appendix_1(text):
    """JMIR wants supplementary tables uploaded as a separate file, with only
    the appendix title in the manuscript. Returns (manuscript, appendix)."""
    start = text.index("# Multimedia Appendix 1")
    end = text.index("# Multimedia Appendix 2")
    body = text[start + len("# Multimedia Appendix 1"):end].strip()
    stub = ("# Multimedia Appendix 1\n\n"
            "Supplementary Tables S1–S7 (uploaded as a separate file).\n\n")
    appendix = "# Multimedia Appendix 1. Supplementary Tables S1–S7\n\n" + body + "\n"
    return text[:start] + stub + text[end:], appendix


def pandoc(text, out):
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False,
                                     encoding="utf-8") as tmp:
        tmp.write(text)
    subprocess.run(["pandoc", tmp.name, "-f", "markdown", "-t", "docx",
                    "--reference-doc", REFERENCE, "--resource-path", DOCS,
                    "-o", out], check=True)
    os.unlink(tmp.name)
    print(f"  ✓ {out}")


def build_metrics_zip():
    """Multimedia Appendix 2: the metric files written by the pipeline scripts
    (stage1_*, stage2_*, tableS*), except per-person out-of-fold predictions. Fixed
    timestamps so an unchanged rebuild gives an identical archive."""
    names = sorted(n for n in os.listdir(OUTPUTS)
                   if n.startswith(("stage1_", "stage2_", "tableS"))
                   and n.endswith((".json", ".csv", ".txt"))
                   and not n.endswith("oof_predictions.csv"))
    with zipfile.ZipFile(APPENDIX_2, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            info = zipfile.ZipInfo(n, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            with open(os.path.join(OUTPUTS, n), "rb") as f:
                z.writestr(info, f.read())
    print(f"  ✓ {APPENDIX_2}  ({len(names)} files)")


if __name__ == "__main__":
    with open(MD, encoding="utf-8") as f:
        text = number_citations(f.read())
    text, appendix = split_appendix_1(text)
    pandoc(text, CLEAN)
    pandoc(appendix, APPENDIX_1)
    with open(APPENDIX_3_MD, encoding="utf-8") as f:
        pandoc(f.read(), APPENDIX_3)
    build_metrics_zip()
    if os.path.exists(V5):
        n, total = highlight_changes(CLEAN, V5, HIGHLIGHTED)
        print(f"  ✓ {HIGHLIGHTED}  ({n}/{total} paragraphs highlighted as new/changed)")
    else:
        print(f"  (skipped highlighted version: {V5} not present)")
